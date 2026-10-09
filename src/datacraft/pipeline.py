"""File adapters, schema rules, SQLite grouping and atomic output."""

from __future__ import annotations

import csv
import json
import logging
import math
import os
import sqlite3
import tempfile
from collections.abc import Iterator
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any

LOG = logging.getLogger(__name__)
Row = dict[str, Any]


class PipelineError(ValueError):
    """An input or configuration violates the pipeline contract."""


def _rows(path: Path) -> Iterator[Row]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise PipelineError("CSV requires unique column names")
            for row in reader:
                if None in row:
                    raise PipelineError("CSV row has more values than its header")
                yield row
    elif suffix == ".jsonl":
        with path.open(encoding="utf-8") as stream:
            for number, line in enumerate(stream, 1):
                if not line.strip():
                    continue
                try:
                    row = json.loads(line, parse_float=Decimal)
                except json.JSONDecodeError as error:
                    raise PipelineError(f"Malformed JSONL at line {number}") from error
                if not isinstance(row, dict):
                    raise PipelineError(f"Expected object at JSONL line {number}")
                yield row
    elif suffix == ".json":
        with path.open(encoding="utf-8") as stream:
            data = json.load(stream, parse_float=Decimal)
        if not isinstance(data, list) or any(not isinstance(row, dict) for row in data):
            raise PipelineError("JSON input must be an array of objects; use JSONL for large files")
        yield from data
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise PipelineError("Install datacraft[excel] to read XLSX") from error
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            header = next(values, ())
            if not header or any(not isinstance(x, str) for x in header):
                raise PipelineError("XLSX first row must contain column names")
            if len(set(header)) != len(header):
                raise PipelineError("XLSX column names must be unique")
            for row in values:
                yield dict(zip(header, row, strict=True))
        finally:
            workbook.close()
    else:
        raise PipelineError(f"Unsupported input format: {suffix}")


def _number(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise PipelineError("Boolean is not a numeric value")
    try:
        text = str(value)
        if len(text) > 1000:
            raise PipelineError("Numeric values are limited to 1000 characters")
        result = Decimal(text)
    except InvalidOperation as error:
        raise PipelineError(f"Invalid decimal: {value!r}") from error
    if not result.is_finite() or abs(result.as_tuple().exponent) > 1000:
        raise PipelineError("Non-finite numbers or exponents above 1000 are rejected")
    return result


def _add_exact(left: Decimal, right: Decimal) -> Decimal:
    # Account for every digit and a carry instead of silently using precision 28.
    precision = (
        max(left.adjusted(), right.adjusted())
        - min(left.as_tuple().exponent, right.as_tuple().exponent)
        + 2
    )
    with localcontext() as context:
        context.prec = max(precision, 28)
        return left + right


def _validate(row: Row, schema: Row) -> Row:
    result = dict(row)
    for name, rule in schema.items():
        value = row.get(name)
        if value is None or value == "":
            if rule.get("required", False):
                raise PipelineError(f"Missing required field: {name}")
            result[name] = None
            continue
        kind = rule.get("type", "str")
        if kind == "decimal":
            value = _number(value)
        elif kind == "int":
            number = _number(value)
            if number != number.to_integral_value():
                raise PipelineError(f"Non-integral value for {name}")
            value = int(number)
        elif kind == "str":
            value = str(value)
        if kind in {"int", "decimal"}:
            if "min" in rule and value < _number(rule["min"]):
                raise PipelineError(f"{name} is below minimum")
            if "max" in rule and value > _number(rule["max"]):
                raise PipelineError(f"{name} is above maximum")
        result[name] = value
    return result


def _matches(row: Row, filters: list[Row]) -> bool:
    for rule in filters:
        name, operator, expected = rule["field"], rule["op"], rule["value"]
        value = row.get(name)
        if operator == "contains":
            match = value is not None and str(expected) in str(value)
        elif operator in {"gt", "ge", "lt", "le"}:
            if value is None:
                return False
            lhs, rhs = _number(value), _number(expected)
            match = {"gt": lhs > rhs, "ge": lhs >= rhs, "lt": lhs < rhs, "le": lhs <= rhs}[operator]
        else:
            match = value == expected or str(value) == str(expected)
            if operator == "ne":
                match = not match
        if not match:
            return False
    return True


def _config(config: Row) -> None:
    if not isinstance(config, dict):
        raise PipelineError("Config must be a JSON object")
    if set(config) - {"schema", "filters", "aggregate", "on_error"}:
        raise PipelineError("Unknown configuration key")
    if config.get("on_error", "fail") not in ("fail", "skip"):
        raise PipelineError("on_error must be fail or skip")
    schema = config.get("schema", {})
    if not isinstance(schema, dict):
        raise PipelineError("schema must be an object")
    for name, rule in schema.items():
        if (
            not isinstance(name, str)
            or not isinstance(rule, dict)
            or rule.get("type", "str") not in ("str", "int", "decimal")
        ):
            raise PipelineError(f"Invalid schema rule: {name}")
        if set(rule) - {"type", "required", "min", "max"}:
            raise PipelineError(f"Unknown schema option for {name}")
        if "required" in rule and not isinstance(rule["required"], bool):
            raise PipelineError("required must be boolean")
        if {"min", "max"} & set(rule) and rule.get("type", "str") == "str":
            raise PipelineError(f"Numeric bounds require int or decimal type: {name}")
        for bound in ("min", "max"):
            if bound in rule:
                _number(rule[bound])
        if "min" in rule and "max" in rule and _number(rule["min"]) > _number(rule["max"]):
            raise PipelineError(f"Minimum exceeds maximum: {name}")
    filters = config.get("filters", [])
    if not isinstance(filters, list):
        raise PipelineError("filters must be an array")
    for rule in filters:
        if (
            not isinstance(rule, dict)
            or set(rule) != {"field", "op", "value"}
            or not isinstance(rule["field"], str)
            or rule["op"] not in ("eq", "ne", "contains", "gt", "ge", "lt", "le")
        ):
            raise PipelineError("Invalid filter rule")
        if rule["op"] in ("gt", "ge", "lt", "le"):
            _number(rule["value"])
    aggregate = config.get("aggregate")
    if aggregate is not None:
        if not isinstance(aggregate, dict) or set(aggregate) != {"group_by", "sum"}:
            raise PipelineError("aggregate requires group_by and sum arrays")
        fields = (
            aggregate["group_by"] + aggregate["sum"]
            if all(isinstance(aggregate[k], list) for k in ("group_by", "sum"))
            else []
        )
        if not fields or any(not isinstance(x, str) for x in fields):
            raise PipelineError("Aggregate field arrays must contain names")
        if len(set(fields)) != len(fields) or "count" in fields:
            raise PipelineError("Aggregate fields must be distinct and must not be count")


def _group_identity(value: Any) -> list[Any]:
    """Tag key types so exact numbers never collide with lookalike strings."""
    if isinstance(value, Decimal):
        number = _number(value)
        with localcontext() as context:
            context.prec = max(len(number.as_tuple().digits), 1)
            canonical = str(number.normalize()) if number else "0"
        return ["decimal", canonical]
    if isinstance(value, list):
        return ["list", [_group_identity(item) for item in value]]
    if isinstance(value, dict):
        return ["object", [[key, _group_identity(item)] for key, item in sorted(value.items())]]
    if isinstance(value, bool):
        return ["bool", value]
    if isinstance(value, int):
        return ["int", value]
    if isinstance(value, float):
        return ["float", _safe_value(value)]
    if value is None:
        return ["null", None]
    return ["str", str(value)]


def _group_value(identity: list[Any]) -> Any:
    kind, value = identity
    if kind == "decimal":
        return Decimal(value)
    if kind == "list":
        return [_group_value(item) for item in value]
    if kind == "object":
        return {key: _group_value(item) for key, item in value}
    return value


def _group(rows: Iterator[Row], group_by: list[str], sums: list[str]) -> Iterator[Row]:
    with tempfile.TemporaryDirectory(prefix="datacraft-") as directory:
        db = sqlite3.connect(str(Path(directory) / "groups.sqlite3"))
        try:
            db.execute("CREATE TABLE groups (key TEXT PRIMARY KEY, n INTEGER, sums TEXT)")
            for row in rows:
                if any(name not in row or row[name] is None for name in group_by + sums):
                    raise PipelineError("Aggregate fields must be present and non-null")
                key = json.dumps([_group_identity(row[x]) for x in group_by], ensure_ascii=False)
                current = db.execute("SELECT n,sums FROM groups WHERE key=?", (key,)).fetchone()
                total = (
                    [Decimal(x) for x in json.loads(current[1])]
                    if current
                    else [Decimal(0)] * len(sums)
                )
                values = [
                    str(_add_exact(a, _number(row[name])))
                    for a, name in zip(total, sums, strict=True)
                ]
                db.execute(
                    "INSERT OR REPLACE INTO groups VALUES (?,?,?)",
                    (key, current[0] + 1 if current else 1, json.dumps(values)),
                )
            db.commit()
            for key, count, total in db.execute("SELECT key,n,sums FROM groups ORDER BY key"):
                yield {
                    **dict(zip(group_by, map(_group_value, json.loads(key)), strict=True)),
                    "count": count,
                    **dict(zip(sums, map(Decimal, json.loads(total)), strict=True)),
                }
        finally:
            db.close()


def _safe_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise PipelineError("Non-finite output value")
    return value


def run(source: Path, destination: Path, config: Row | None = None) -> dict[str, int]:
    """Apply a pipeline atomically. Decimal values export as strings without float loss.

    Source and destination must differ. Existing output survives all processing errors.
    Skip policy applies only to schema/filter validation; malformed files fail the pipeline.
    """
    config = {} if config is None else config
    _config(config)
    if source.resolve() == destination.resolve():
        raise PipelineError("Source and destination must differ")
    if destination.suffix.lower() not in {".csv", ".json", ".jsonl"}:
        raise PipelineError("Output format must be CSV, JSON or JSONL")
    stats = {"read": 0, "accepted": 0, "rejected": 0, "filtered": 0, "written": 0}

    def selected() -> Iterator[Row]:
        for number, row in enumerate(_rows(source), 1):
            stats["read"] += 1
            try:
                item = _validate(row, config.get("schema", {}))
                matches = _matches(item, config.get("filters", []))
            except PipelineError as error:
                if config.get("on_error", "fail") == "fail":
                    raise PipelineError(f"Row {number}: {error}") from error
                stats["rejected"] += 1
                LOG.warning(
                    json.dumps({"event": "row_rejected", "row": number, "reason": str(error)})
                )
                continue
            if not matches:
                stats["filtered"] += 1
                continue
            stats["accepted"] += 1
            yield item

    rows = selected()
    if config.get("aggregate") is not None:
        aggregate = config["aggregate"]
        rows = _group(rows, aggregate["group_by"], aggregate["sum"])
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".datacraft-", dir=destination.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="") as stream:
            writer = None
            fields: list[str] | None = None
            if destination.suffix.lower() == ".json":
                stream.write("[\n")
            for row in rows:
                row = {key: _safe_value(value) for key, value in row.items()}
                if destination.suffix.lower() == ".csv":
                    if writer is None:
                        fields = list(row)
                        writer = csv.DictWriter(stream, fieldnames=fields)
                        writer.writeheader()
                    if set(row) != set(fields or []):
                        raise PipelineError("CSV output requires a consistent schema")
                    # Prevent spreadsheet formula evaluation in exported untrusted text.
                    row = {
                        k: "'" + v
                        if isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
                        else v
                        for k, v in row.items()
                    }
                    writer.writerow(row)
                else:
                    if destination.suffix.lower() == ".json" and stats["written"]:
                        stream.write(",\n")
                    json.dump(row, stream, ensure_ascii=False, allow_nan=False, default=str)
                    if destination.suffix.lower() == ".jsonl":
                        stream.write("\n")
                stats["written"] += 1
            if destination.suffix.lower() == ".json":
                stream.write("\n]\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        raise
    LOG.info(json.dumps({"event": "pipeline_complete", **stats}))
    return stats
